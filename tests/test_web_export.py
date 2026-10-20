"""The export page and the download."""

import io
import json
import zipfile

import pytest

pytest.importorskip("fastapi")
pytest.importorskip("multipart")

from openpyxl import load_workbook
from support.study import line_id, make_line, make_script
from support.web import ORIGIN, make_client

from yomineko.pipeline.workspace import Workspace
from yomineko.script_io import write_script
from yomineko.web.app import KEPT_DOWNLOADS


@pytest.fixture
def workspace(tmp_path):
    return Workspace(tmp_path / "games")


@pytest.fixture
def client(workspace):
    write_script(make_script(), workspace.script_path("demo"))
    return make_client(workspace)


def download(client, **data):
    return client.post("/games/demo/export", data=data, headers=ORIGIN, follow_redirects=True)


def attachment(response):
    return response.headers["content-disposition"]


# ----------------------------------------------------------------------- page


def test_the_form_offers_the_formats_and_options(client):
    page = client.get("/games/demo/export")
    assert page.status_code == 200
    for text in ("Anki deck", "spreadsheet", "CSV", "JSON", "More options", "Download"):
        assert text in page.text
    assert 'value="ハル"' in page.text and 'name="decks"' in page.text


def test_the_form_of_an_unanalysed_game_says_what_still_works(workspace):
    write_script(make_script(words=[]), workspace.script_path("demo"))
    page = make_client(workspace).get("/games/demo/export")
    assert "Find the words first" in page.text


def test_the_game_page_links_to_the_export(client):
    assert "/games/demo/export" in client.get("/games/demo").text


def test_the_export_page_of_a_missing_game_is_a_clean_error(workspace):
    page = make_client(workspace).get("/games/demo/export")
    assert page.status_code == 400 and "not been extracted" in page.text


# ------------------------------------------------------------------- downloads


def test_an_anki_deck_is_downloaded(client):
    response = download(client, kind="anki", decks=["words", "sentences", "terms"])
    assert response.status_code == 200
    assert attachment(response) == 'attachment; filename="demo.apkg"'
    assert zipfile.is_zipfile(io.BytesIO(response.content))


def test_the_form_post_sends_the_page_to_the_file_by_get(client):
    response = client.post(
        "/games/demo/export", data={"kind": "csv"}, headers=ORIGIN, follow_redirects=False
    )
    assert response.status_code == 303
    assert response.headers["location"].startswith("/downloads/")


def test_the_file_can_be_fetched_again(client):
    # the macOS window saves a download by fetching its address a second time
    first = download(client, kind="anki", decks=["words"])
    again = client.get(first.url)
    assert again.status_code == 200 and again.content == first.content
    assert attachment(again) == 'attachment; filename="demo.apkg"'


@pytest.mark.parametrize(
    ("data", "name"),
    [
        ({"kind": "csv"}, "demo-words.csv"),
        ({"kind": "json"}, "demo-words.json"),
        ({"kind": "xlsx", "decks": ["words"]}, "demo.xlsx"),
        ({"kind": "anki", "decks": ["words"]}, "demo.apkg"),
    ],
)
def test_every_export_is_a_type_a_window_cannot_show(client, data, name):
    # WebKit shows text/csv and application/json instead of saving them
    response = download(client, **data)
    assert response.headers["content-type"] == "application/octet-stream"
    assert attachment(response) == f'attachment; filename="{name}"'


def test_only_the_newest_files_are_kept(client):
    urls = [download(client, kind="csv", min_freq=str(n)).url for n in range(1, KEPT_DOWNLOADS + 2)]
    gone = client.get(urls[0])
    assert gone.status_code == 400 and "no longer kept" in gone.text
    assert all(client.get(url).status_code == 200 for url in urls[1:])


def test_an_unknown_download_is_a_clean_error(client):
    response = client.get("/downloads/guessed")
    assert response.status_code == 400 and "no longer kept" in response.text


def test_a_spreadsheet_is_downloaded(client):
    response = download(client, kind="xlsx", decks=["words"])
    assert attachment(response) == 'attachment; filename="demo.xlsx"'
    book = load_workbook(io.BytesIO(response.content))
    assert len(list(book["Words"].iter_rows())) == 61


def test_a_word_list_csv_has_every_word_by_default(client):
    response = download(client, kind="csv", content="words")
    assert attachment(response) == 'attachment; filename="demo-words.csv"'
    assert response.content.startswith(b"\xef\xbb\xbf") and response.text.count("\n") >= 61


def test_sentences_as_json_honour_the_count_and_order(client):
    response = download(client, kind="json", content="sentences", sentences="2", order="rom")
    assert attachment(response) == 'attachment; filename="demo-sentences.json"'
    assert len(json.loads(response.content)["lines"]) == 2


def test_lines_as_csv(client):
    response = download(client, kind="csv", content="lines")
    assert attachment(response) == 'attachment; filename="demo-lines.csv"'


def test_more_options_are_applied(client):
    response = download(
        client, kind="xlsx", decks=["words"], min_freq="950", player_name="", deck_name=""
    )
    book = load_workbook(io.BytesIO(response.content))
    assert len(list(book["Words"].iter_rows())) == 51  # 50 words and the header


def test_the_players_name_is_used_on_cards(workspace):
    lines = [make_line(1, "〇〇は　あたらしいまちへ。", [1, 2])]
    write_script(make_script(lines=lines), workspace.script_path("demo"))
    response = download(
        make_client(workspace), kind="xlsx", decks=["sentences"], player_name="タロウ"
    )
    sheet = load_workbook(io.BytesIO(response.content))["Sentences"]
    assert sheet["B2"].value == "タロウは　あたらしいまちへ。"


def test_the_users_review_decides_the_sentences(workspace):
    write_script(make_script(), workspace.script_path("demo"))
    workspace.review_path("demo").write_text(
        f'sentences:\n  exclude: ["{line_id(1)}"]\n', encoding="utf-8"
    )
    response = download(make_client(workspace), kind="json", content="sentences")
    assert len(json.loads(response.content)["lines"]) == 4


# --------------------------------------------------------------------- mistakes


@pytest.mark.parametrize(
    ("data", "message"),
    [
        ({"kind": "pdf"}, "Choose what to export"),
        ({"kind": "csv", "content": "lyrics"}, "Choose what to export"),
        ({"kind": "csv", "order": "random"}, "Choose what to export"),
        ({"kind": "anki"}, "Choose at least one"),
        ({"kind": "anki", "decks": ["cards"]}, "unknown deck"),
        ({"kind": "csv", "min_freq": "many"}, "is not a whole number"),
        ({"kind": "csv", "sentences": "-3"}, "is not a whole number"),
    ],
)
def test_mistakes_are_explained(client, data, message):
    response = download(client, **data)
    assert response.status_code == 400 and message in response.text


def test_a_game_that_is_not_analysed_cannot_be_exported_as_a_deck(workspace):
    write_script(make_script(words=[]), workspace.script_path("demo"))
    response = download(make_client(workspace), kind="anki", decks=["words"])
    assert response.status_code == 400 and "yomineko analyze" in response.text


def test_an_export_request_from_another_site_is_refused(client):
    response = client.post("/games/demo/export", data={"kind": "csv"}, follow_redirects=False)
    assert response.status_code == 403
