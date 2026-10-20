from openpyxl import load_workbook
from support.study import CREDIT, make_line, make_script

from yomineko.export.sentences import select_sentences
from yomineko.export.study import build_study
from yomineko.export.xlsx_export import SENTENCE_HEADERS, TERM_HEADERS, WORD_HEADERS, write_xlsx


def study_of(script=None, **kw):
    script = script or make_script()
    return build_study(script, select_sentences(script), min_freq=4, **kw)


def rows(book, name):
    return [list(r) for r in book[name].iter_rows(values_only=True)]


def test_the_workbook_has_a_sheet_each_and_an_about_page(tmp_path):
    dest = tmp_path / "out" / "deck.xlsx"
    count = write_xlsx(study_of(), dest)
    book = load_workbook(dest)
    assert book.sheetnames == ["Words", "Sentences", "Terms", "About"]
    assert count == 60 + 5 + 1
    assert tuple(rows(book, "Words")[0]) == WORD_HEADERS
    assert tuple(rows(book, "Sentences")[0]) == SENTENCE_HEADERS
    assert tuple(rows(book, "Terms")[0]) == TERM_HEADERS


def test_rows_hold_the_card_content(tmp_path):
    dest = tmp_path / "d.xlsx"
    write_xlsx(study_of(), dest)
    book = load_workbook(dest)
    first = rows(book, "Words")[1]
    assert first[:6] == [1, "こと", "事", "こと", "名詞", 999]
    assert first[6] == "meaning 1; extra 1 | second 1" and first[8] == 1
    sentence = next(r for r in rows(book, "Sentences") if r[0] == "08000003")
    assert sentence[1] == "そのほんは　とてもおもしろい。" and "ことば44: meaning 44" in sentence[4]
    assert rows(book, "Terms")[1] == ["ボタン", "item", "button", 900]


def test_about_carries_the_counts_and_the_credit(tmp_path):
    dest = tmp_path / "d.xlsx"
    write_xlsx(study_of(), dest)
    about = dict(rows(load_workbook(dest), "About"))
    assert about["game"] == "Demo Game" and about["words"] == 60 and about["sentences"] == 5
    assert about["dictionary"] == CREDIT


def test_text_that_looks_like_a_formula_stays_text(tmp_path):
    script = make_script()
    script = script.model_copy(
        update={"lines": [make_line(1, "=1+1　ようこそあたらしいまちへ。", [2, 50])]}
    )
    dest = tmp_path / "f.xlsx"
    write_xlsx(study_of(script), dest)
    cell = load_workbook(dest)["Sentences"]["B2"]
    assert cell.data_type == "s" and cell.value.startswith("=1+1")


def test_a_deck_left_out_is_an_empty_sheet(tmp_path):
    dest = tmp_path / "d.xlsx"
    write_xlsx(study_of(decks=["words"]), dest)
    book = load_workbook(dest)
    assert len(rows(book, "Sentences")) == 1 and len(rows(book, "Terms")) == 1
