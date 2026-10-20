"""``yomineko export`` with sentences, Anki and XLSX."""

import csv
import io
import json
import re
import zipfile

import pytest
from openpyxl import load_workbook
from support.study import line_id, make_line, make_script
from typer.testing import CliRunner

from yomineko.cli.app import app
from yomineko.script_io import read_script, write_script

runner = CliRunner()
ANSI = re.compile(chr(27) + r"\[[0-9;]*m")


def run(*args):
    return runner.invoke(app, [str(a) for a in args])


def plain(result):
    return ANSI.sub("", result.output)


@pytest.fixture
def script_file(tmp_path):
    return write_script(make_script(), tmp_path / "script.json")


def csv_rows(path):
    text = path.read_text(encoding="utf-8-sig")
    return [r for r in csv.reader(io.StringIO(text)) if r and not r[0].startswith("#")]


def test_sentences_csv_lists_the_picks_in_easy_to_hard_order(script_file):
    out = script_file.with_name("s.csv")
    result = run("export", script_file, "-f", "csv", "--content", "sentences", "-o", out)
    assert result.exit_code == 0, result.output
    assert "Exported 5 sentences" in result.output
    rows = csv_rows(out)
    assert rows[0] == ["id", "text", "english", "tier", "words"]
    assert len(rows) == 6 and rows[1][4] == "1 2"  # the shortest comes first inside tier 0
    assert "CC BY-SA" in out.read_text(encoding="utf-8-sig")


def test_sentences_in_rom_order_and_limited(script_file):
    out = script_file.with_name("s.csv")
    result = run(
        "export", script_file, "-f", "csv", "--content", "sentences", "--order", "rom",
        "--sentences", "3", "-o", out,
    )  # fmt: skip
    assert result.exit_code == 0, result.output
    ids = [r[0] for r in csv_rows(out)[1:]]
    assert len(ids) == 3 and ids == sorted(ids)


def test_sentences_json_is_a_valid_script_with_the_picked_lines(script_file):
    out = script_file.with_name("s.json")
    result = run("export", script_file, "-f", "json", "--content", "sentences", "-o", out)
    assert result.exit_code == 0, result.output
    narrowed = read_script(out)
    assert len(narrowed.lines) == 5 and len(narrowed.words) == 60


def test_the_review_can_exclude_and_pin_lines(script_file, tmp_path):
    review = tmp_path / "review.yaml"
    review.write_text(
        f'sentences:\n  exclude: ["{line_id(1)}"]\n  pin: ["{line_id(2)}"]\n', encoding="utf-8"
    )
    out = tmp_path / "s.csv"
    result = run(
        "export", script_file, "-f", "csv", "--content", "sentences", "-r", review, "-o", out
    )
    assert result.exit_code == 0, result.output
    ids = {r[0] for r in csv_rows(out)[1:]}
    assert line_id(1) not in ids and line_id(2) in ids and len(ids) == 4


def test_a_packaged_profile_adds_its_ranges(tmp_path):
    from yomineko.profile.loader import packaged_profiles

    pmd = next(p for p in packaged_profiles() if p.profile.id == "pmd-red-jp")
    assert pmd.profile.sentences.exclude_ranges and pmd.profile.sentences.caps.name_quote == 100
    game = make_script().game.model_copy(update={"profile_id": "pmd-red-jp"})
    lines = [
        make_line(0x30A000 + 4, "ようこそ　あたらしいまちへ。", [1, 2]),
        make_line(0x20, "きょうは　いいてんきですね。", [2, 50]),
    ]
    script = write_script(make_script(game=game, lines=lines), tmp_path / "p.json")
    out = tmp_path / "s.csv"
    result = run(
        "export", script, "-f", "csv", "--content", "sentences", "--no-packaged-review", "-o", out
    )
    assert result.exit_code == 0, result.output
    assert [r[1] for r in csv_rows(out)[1:]] == ["きょうは　いいてんきですね。"]


def test_an_unknown_profile_file_is_an_error(script_file):
    result = run("export", script_file, "-f", "csv", "--content", "sentences", "-p", "nope.yaml")
    assert result.exit_code == 1 and "error:" in result.output


def test_anki_deck_is_written_next_to_the_script(script_file):
    result = run("export", script_file, "-f", "anki")
    assert result.exit_code == 0, result.output
    deck = script_file.with_name("script.export.apkg")
    assert deck.is_file() and zipfile.is_zipfile(deck)
    assert "Wrote 66 cards (60 words, 5 sentences, 1 terms)" in result.output


def test_anki_decks_and_name_options(script_file, tmp_path):
    out = tmp_path / "x.apkg"
    result = run(
        "export", script_file, "-f", "anki", "--decks", "terms,words", "--deck-name", "D",
        "--min-freq", "999", "-o", out,
    )  # fmt: skip
    assert result.exit_code == 0, result.output
    assert "(1 words, 0 sentences, 1 terms)" in result.output  # only id 1 has freq 999


def test_xlsx_workbook_is_written(script_file, tmp_path):
    out = tmp_path / "x.xlsx"
    assert run("export", script_file, "-f", "xlsx", "-o", out).exit_code == 0
    assert load_workbook(out).sheetnames == ["Words", "Sentences", "Terms", "About"]


@pytest.mark.parametrize(
    ("args", "message"),
    [
        (["-f", "anki", "--content", "words"], "--content does not apply"),
        (["-f", "csv", "--decks", "words"], "--decks applies to anki and xlsx"),
        (["-f", "anki", "--decks", "cards"], "unknown deck"),
        (["-f", "csv", "--content", "sentences", "--all"], "--all and --kind"),
    ],
)
def test_option_mistakes_are_clear(script_file, args, message):
    result = run("export", script_file, *args)
    assert result.exit_code == 1 and message in plain(result)


def test_an_unanalysed_script_says_to_run_analyze(tmp_path):
    path = write_script(make_script(words=[]), tmp_path / "raw.json")
    for fmt in ("anki", "xlsx"):
        result = run("export", path, "-f", fmt)
        assert result.exit_code == 1 and "yomineko analyze" in result.output


def test_a_script_analysed_before_schema_3_says_to_analyze_again(tmp_path):
    lines = [make_line(1, "きょうは　いいてんきですね。")]  # no word ids
    path = write_script(make_script(lines=lines), tmp_path / "old.json")
    result = run("export", path, "-f", "anki")
    assert result.exit_code == 1 and "again" in result.output


def test_the_word_list_still_defaults_to_every_word(script_file):
    out = script_file.with_name("w.json")
    result = run("export", script_file, "-f", "json", "--content", "words", "-o", out)
    assert result.exit_code == 0
    assert len(json.loads(out.read_text(encoding="utf-8"))["words"]) == 60


def test_player_name_option_sets_the_stand_in(tmp_path):
    lines = [
        make_line(1, "〇〇は　あたらしいまちへ。", [1, 2]),
        make_line(2, "きょうは　いいてんきですね。", [2, 50]),
    ]
    script = write_script(make_script(lines=lines), tmp_path / "n.json")
    out = tmp_path / "n.xlsx"
    result = run("export", script, "-f", "xlsx", "--player-name", "タロウ", "-o", out)
    assert result.exit_code == 0, result.output
    texts = [r[1] for r in load_workbook(out)["Sentences"].iter_rows(min_row=2, values_only=True)]
    assert "タロウは　あたらしいまちへ。" in texts
    csv_out = tmp_path / "n.csv"
    run("export", script, "-f", "csv", "--content", "sentences", "-o", csv_out)
    assert "〇〇は" in csv_out.read_text(encoding="utf-8-sig"), "data exports keep the placeholder"
