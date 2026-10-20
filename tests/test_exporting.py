"""The shared export builder behind the CLI and the web download."""

import io
import json
import zipfile

import pytest
from openpyxl import load_workbook
from support.study import line_id, make_line, make_script

from yomineko.errors import YominekoError
from yomineko.export.words_export import NotAnalysedError
from yomineko.pipeline.exporting import ExportRequest, build_export


def export(**kw):
    kw.setdefault("format", "csv")
    return build_export(make_script(), ExportRequest(**kw))


def test_an_anki_deck_is_a_zip_with_every_card():
    made = export(format="anki")
    assert made.extension == "apkg" and made.media_type == "application/octet-stream"
    assert zipfile.is_zipfile(io.BytesIO(made.data))
    assert made.summary == "Wrote 66 cards (60 words, 5 sentences, 1 terms)"


def test_a_spreadsheet_opens_and_has_the_chosen_parts():
    made = export(format="xlsx", decks=["terms"])
    assert made.extension == "xlsx" and "spreadsheetml" in made.media_type
    book = load_workbook(io.BytesIO(made.data))
    assert len(list(book["Terms"].iter_rows())) == 2 and len(list(book["Words"].iter_rows())) == 1
    assert made.summary == "Wrote 1 cards (0 words, 0 sentences, 1 terms)"


def test_csv_is_written_with_a_byte_order_mark_for_excel():
    made = export(format="csv", content="words")
    assert made.data.startswith(b"\xef\xbb\xbf") and made.extension == "csv"
    assert made.summary == "Exported 60 words"


def test_json_is_plain_utf8():
    made = export(format="json", content="lines")
    assert not made.data.startswith(b"\xef\xbb\xbf") and made.media_type == "application/json"
    assert len(json.loads(made.data)["lines"]) == 5


@pytest.mark.parametrize("fmt", ["csv", "json"])
def test_sentences_follow_the_order_and_count(fmt):
    easy = export(format=fmt, content="sentences", sentences=3)
    assert easy.summary == "Exported 3 sentences"
    rom = export(format=fmt, content="sentences", order="rom")
    assert rom.summary == "Exported 5 sentences"


def test_the_minimum_frequency_defaults_differ_for_words_and_decks():
    assert export(format="csv", content="words").summary == "Exported 60 words"
    assert export(format="csv", content="words", min_freq=950).summary == "Exported 50 words"
    # only 3 sentences keep two words that are seen that often
    assert export(format="anki", min_freq=950).summary == (
        "Wrote 54 cards (50 words, 3 sentences, 1 terms)"
    )


def test_the_players_name_and_deck_name_are_used():
    lines = [make_line(1, "〇〇は　あたらしいまちへ。", [1, 2])]
    script = make_script(lines=lines)
    made = build_export(
        script, ExportRequest(format="xlsx", player_name="タロウ", decks=["sentences"])
    )
    sheet = load_workbook(io.BytesIO(made.data))["Sentences"]
    assert sheet["B2"].value == "タロウは　あたらしいまちへ。"


def test_review_files_decide_which_sentences_are_picked(tmp_path):
    review = tmp_path / "r.yaml"
    review.write_text(f'sentences:\n  exclude: ["{line_id(1)}"]\n', encoding="utf-8")
    made = export(format="csv", content="sentences", review_files=[review])
    assert made.summary == "Exported 4 sentences"


def test_lines_can_be_filtered_by_kind():
    lines = [make_line(1, "ようこそ　あたらしいまちへ。"), make_line(2, "ボタン", kind="term")]
    script = make_script(lines=lines)
    made = build_export(script, ExportRequest(format="csv", content="lines", kind="term"))
    assert made.summary == "Exported 1 lines"


def test_a_script_that_is_not_analysed_is_refused_with_the_next_step():
    script = make_script(words=[])
    with pytest.raises(NotAnalysedError, match="yomineko analyze"):
        build_export(script, ExportRequest(format="anki"))
    with pytest.raises(NotAnalysedError):
        build_export(script, ExportRequest(format="csv", content="words"))
    assert build_export(script, ExportRequest(format="csv", content="lines")).summary


def test_an_unknown_content_is_refused():
    with pytest.raises(YominekoError, match="unknown content"):
        export(format="csv", content="everything")  # type: ignore[arg-type]


def test_the_file_name_is_made_from_the_game_and_the_content():
    made = export(format="csv", content="words")
    assert made.filename("pmd-red-jp", "words") == "pmd-red-jp-words.csv"
    assert made.filename("pmd-red-jp") == "pmd-red-jp.csv"
