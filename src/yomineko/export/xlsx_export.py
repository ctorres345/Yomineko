"""Write the study material as a spreadsheet (.xlsx): Words, Sentences, Terms and About sheets.

Cells are always written as text or numbers, never as formulas, so a line that happens to start
with ``=`` stays text.
"""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from yomineko.export.study import StudySet

WORD_HEADERS = (
    "rank",
    "word",
    "kanji",
    "reading",
    "pos",
    "freq",
    "meanings",
    "example",
    "jmdict_id",
)
SENTENCE_HEADERS = ("id", "japanese", "english", "tier", "words")
TERM_HEADERS = ("japanese", "category", "meaning", "jmdict_id")


def _append(sheet: Worksheet, row: Sequence[str | int]) -> None:
    sheet.append(list(row))
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"  # never a formula


def _header(sheet: Worksheet, headers: Sequence[str], widths: Sequence[int]) -> None:
    _append(sheet, headers)
    sheet.freeze_panes = "A2"
    for letter, width in zip("ABCDEFGHIJ", widths, strict=False):
        sheet.column_dimensions[letter].width = width


def write_xlsx(study: StudySet, dest: str | Path) -> int:
    """Write the workbook; returns the number of rows across the three data sheets."""
    book = Workbook()
    words = book.active
    assert words is not None
    words.title = "Words"
    _header(words, WORD_HEADERS, (7, 16, 12, 16, 10, 7, 50, 50, 11))
    for w in study.words:
        _append(
            words,
            [
                w.rank,
                w.front,
                w.kanji,
                w.kana,
                w.pos,
                w.freq,
                " | ".join(w.meanings),
                w.example,
                w.jmdict_id,
            ],
        )
    sentences = book.create_sheet("Sentences")
    _header(sentences, SENTENCE_HEADERS, (11, 60, 40, 6, 60))
    for s in study.sentences:
        listed = "; ".join(
            f"{g.kana}{f' ({g.kanji})' if g.kanji else ''}: {g.meaning}" for g in s.words
        )
        _append(sentences, [s.line_id, s.text, s.english, s.tier, listed])
    terms = book.create_sheet("Terms")
    _header(terms, TERM_HEADERS, (24, 14, 40, 11))
    for t in study.terms:
        _append(terms, [t.jp, t.category, t.meaning, t.jmdict_id])
    about = book.create_sheet("About")
    for row in (
        ("game", study.title),
        ("profile", study.profile_id),
        ("words", len(study.words)),
        ("sentences", len(study.sentences)),
        ("terms", len(study.terms)),
        ("minimum word frequency", study.min_freq),
        ("dictionary", study.credit),
    ):
        _append(about, row)
    about.column_dimensions["A"].width = 24
    path = Path(dest)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return len(study.words) + len(study.sentences) + len(study.terms)
